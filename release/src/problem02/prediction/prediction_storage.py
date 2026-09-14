"""生成 2025 年 1 月每天 0:00 的储能电量表。

运行：
    python result/problem02/prediction/prediction_storage.py

输出：
    result/problem02/prediction/january_soc_2025.xlsx
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from src.problem02 import problem

ROOT = Path(__file__).resolve().parents[3]
RESULT_DIR = ROOT / "result" / "problem02" / "prediction"
APPENDIX_DIR = ROOT / "附件"

def main() -> None:
    attachment = APPENDIX_DIR / "附件2.xlsx"
    price = problem.load_price(
        attachment.parent / "附件1.xlsx"
    )
    actual_dates, actual_load, actual_pv = problem.load_actual_series(attachment)
    all_dates, all_load_pred, all_pv_pred = problem.build_january_forecasts(
        attachment
    )
    jan_indices = [
        index for index, target in enumerate(all_dates) if target.month == 1
    ]
    jan_dates = [all_dates[index] for index in jan_indices]
    jan_load_pred = all_load_pred[jan_indices]
    jan_pv_pred = all_pv_pred[jan_indices]
    if jan_dates[0] != date(2025, 1, 1) or jan_dates[-1] != date(2025, 1, 31):
        raise ValueError("1 月预测日期范围不正确")

    actual_by_date = {
        target: index for index, target in enumerate(actual_dates)
    }
    jan_actual_indices = [actual_by_date[target] for target in jan_dates]
    actual_net = (
        actual_load[jan_actual_indices] - actual_pv[jan_actual_indices]
    ) * problem.DT
    predicted_net = (jan_load_pred - jan_pv_pred) * problem.DT
    residual_history = list(actual_net - predicted_net)

    storage = problem.SOC_INITIAL
    rows = []
    for index, target in enumerate(jan_dates):
        rows.append(
            {
                "日期": target,
                "0:00储电量（kWh）": storage,
            }
        )

        forecast_net = predicted_net[index]
        recent = np.asarray(
            residual_history[-problem.RESIDUAL_WINDOW_DAYS:], dtype=float
        )
        risk_margin = np.quantile(
            recent,
            problem.RISK_QUANTILE,
            axis=0,
            method="linear",
        )
        plan = problem.solve_plan_day(
            forecast_net + risk_margin,
            price,
            storage,
        )
        actual = problem.simulate_actual_day(
            plan.purchase_kwh,
            actual_load[actual_by_date[target]],
            actual_pv[actual_by_date[target]],
            storage,
        )
        storage = float(actual.storage_end_kwh[-1])

    output = RESULT_DIR / "january_soc_2025.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "1月0点储电量"
    worksheet.append(["日期", "0:00储电量（kWh）"])
    for row in rows:
        worksheet.append(
            [
                datetime.combine(row["日期"], datetime.min.time()),
                float(row["0:00储电量（kWh）"]),
            ]
        )

    for cell in worksheet[1]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(name="宋体", size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        row[0].alignment = Alignment(horizontal="center")
        row[1].number_format = "0.000000"
    worksheet.column_dimensions["A"].width = 14
    worksheet.column_dimensions["B"].width = 20
    worksheet.freeze_panes = "A2"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    print(f"2025-01-01 0:00 储电量：{rows[0]['0:00储电量（kWh）']:.6f} kWh")
    print(f"2025-01-31 0:00 储电量：{rows[-1]['0:00储电量（kWh）']:.6f} kWh")
    print(f"2025-02-01 0:00 储电量：{storage:.6f} kWh")
    print(f"表格已保存：{output}")


if __name__ == "__main__":
    main()
