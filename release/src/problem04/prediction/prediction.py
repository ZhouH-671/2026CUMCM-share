"""第四问正式滚动均值电价预测模型。

本脚本可以独立运行，不导入第四问其他 Python 文件，只需要 numpy + openpyxl。

模型
----
1) 目标序列：附件4 的**日均价**  y(d) = mean(附件4 第 d 天 144 个时段)
2) 预测：滚动均值（无泄露，只用 d 之前的数据）
       y_hat(d) = mean( y(d-K), ..., y(d-1) )        K = 7（--window）
3) 重建到逐时段
       偏离水平       = y_hat(d) - mean(附件1)
       pred(d, slot)  = max( 附件1(slot) + 偏离水平, 0 )

即"附件1 的经验分时形状 + 滚动均值给出的价格水平"。
附件1 只作为基准形状使用，不参与统计，因此不存在信息泄露。

输入
----
附件/附件1.xlsx    分时电价基准形状
附件/附件4.xlsx    实际电价，365 天 x 144 时段

输出
----
result/problem04/prediction/pred_price_rolling.xlsx
    单张工作表「预测电价」：334 行（2025-02-01 ~ 2025-12-31）× 144 列
    首列日期，首行时段表头（取自附件4）

运行（在项目根目录下）
----
    python -m src.problem04.prediction.prediction
    python -m src.problem04.prediction.prediction --window 14
    python -m src.problem04.prediction.prediction --output result/problem04/my_rolling.xlsx
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
from pathlib import Path

import numpy as np
import openpyxl

PERIODS_PER_DAY = 144                      # 一天 144 个 10 分钟时段
WINDOW = 7                                 # 滚动均值窗口（天）
OUTPUT_START = date(2025, 2, 1)            # 评估/输出区间起点
OUTPUT_END = date(2025, 12, 31)            # 评估/输出区间终点
ATTACHMENT1 = "附件1.xlsx"
ATTACHMENT4 = "附件4.xlsx"

ROOT = Path(__file__).resolve().parents[3]
APPENDIX_DIR = ROOT / "附件"
RESULT_DIR = ROOT / "result" / "problem04" / "prediction"


# --------------------------------------------------------------------------
# 读取 Excel（本文件内自带，不依赖任何外部模块）
# --------------------------------------------------------------------------
def as_date(value) -> date:
    """把 Excel 里的日期单元格转成 datetime.date。"""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise TypeError("日期格式错误：%r" % (value,))



def load_baseline_shape(path: Path) -> np.ndarray:
    """附件1：144 个时段的分时电价（经验形状）。"""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    values = np.asarray(
        [row[1] for row in rows[1:] if row[0] is not None], dtype=float
    )
    if values.shape != (PERIODS_PER_DAY,):
        raise ValueError("附件1 电价列应为 %d 个数值，实际 %s"
                         % (PERIODS_PER_DAY, values.shape))
    return values


def load_actual_price(path: Path):
    """附件4：返回 (日期列表, 365 x 144 电价矩阵, 144 个时段表头)。"""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    header = list(rows[0][1:1 + PERIODS_PER_DAY])
    dates, prices = [], []
    for row in rows[1:]:
        if row[0] is None:
            continue
        dates.append(as_date(row[0]))
        prices.append([float(value) for value in row[1:1 + PERIODS_PER_DAY]])
    return dates, np.asarray(prices, dtype=float), header


# --------------------------------------------------------------------------
# 滚动均值
# --------------------------------------------------------------------------
def rolling_levels(daily: np.ndarray, window: int) -> np.ndarray:
    """第 i 天的水平预测 = 前 window 天日均价的均值（i=0 没有历史，记 nan）。"""
    levels = np.full(daily.shape, np.nan, dtype=float)
    for index in range(1, len(daily)):
        levels[index] = float(daily[max(0, index - window):index].mean())
    return levels


def reconstruct(level: float, baseline: np.ndarray, baseline_mean: float) -> np.ndarray:
    """把"日均价水平"还原成 144 个时段的电价。"""
    return np.maximum(baseline + (level - baseline_mean), 0.0)


def evaluate(predicted: np.ndarray, actual: np.ndarray) -> dict:
    error = predicted - actual
    ss_res = float((error ** 2).sum())
    ss_tot = float(((actual - actual.mean()) ** 2).sum())
    pred_daily, actual_daily = predicted.mean(axis=1), actual.mean(axis=1)
    ss_res_daily = float(((pred_daily - actual_daily) ** 2).sum())
    ss_tot_daily = float(((actual_daily - actual_daily.mean()) ** 2).sum())
    return {
        "mae_slot": float(np.abs(error).mean()),
        "rmse_slot": float(np.sqrt((error ** 2).mean())),
        "r2_slot": 1.0 - ss_res / ss_tot if ss_tot > 0 else float("nan"),
        "mae_daily": float(np.abs(pred_daily - actual_daily).mean()),
        "rmse_daily": float(np.sqrt(((pred_daily - actual_daily) ** 2).mean())),
        "r2_daily": (1.0 - ss_res_daily / ss_tot_daily
                     if ss_tot_daily > 0 else float("nan")),
    }


# --------------------------------------------------------------------------
# 输出
# --------------------------------------------------------------------------
def write_workbook(path: Path, header: list, dates: list,
                   predicted: np.ndarray) -> None:
    """单张「预测电价」工作表：首列日期，其余 144 列为各时段预测电价。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "预测电价"
    sheet.append(["日期"] + list(header))
    for index, day in enumerate(dates):
        sheet.append([datetime.combine(day, datetime.min.time())]
                     + [float(value) for value in predicted[index]])
    sheet.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--window", type=int, default=WINDOW,
                        help="滚动均值窗口天数（默认 %d）" % WINDOW)
    parser.add_argument("--start", type=str, default=str(OUTPUT_START))
    parser.add_argument("--end", type=str, default=str(OUTPUT_END))
    parser.add_argument("--output", type=Path,
                        default=RESULT_DIR / "pred_price_rolling.xlsx")
    args = parser.parse_args()

    start = datetime.strptime(args.start, "%Y-%m-%d").date()
    end = datetime.strptime(args.end, "%Y-%m-%d").date()

    baseline = load_baseline_shape(APPENDIX_DIR / ATTACHMENT1)
    price_dates, price_matrix, header = load_actual_price(APPENDIX_DIR / ATTACHMENT4)
    print("附件目录：%s" % APPENDIX_DIR)
    print("附件4：%d 天 x %d 时段，日均价 %.4f 元/kWh"
          % (price_matrix.shape[0], price_matrix.shape[1], price_matrix.mean()))
    print("附件1 基准形状：均值 %.4f 元/kWh" % baseline.mean())

    daily = price_matrix.mean(axis=1)
    levels = rolling_levels(daily, args.window)
    baseline_mean = float(baseline.mean())

    eval_dates = [day for day in price_dates if start <= day <= end]
    predicted = np.stack([
        reconstruct(levels[price_dates.index(day)], baseline, baseline_mean)
        if np.isfinite(levels[price_dates.index(day)]) else baseline.copy()
        for day in eval_dates
    ])
    actual = np.stack([price_matrix[price_dates.index(day)] for day in eval_dates])

    metrics = evaluate(predicted, actual)
    print("评估区间：%s ~ %s（%d 天），滚动窗口 K = %d"
          % (eval_dates[0], eval_dates[-1], len(eval_dates), args.window))
    print("预测精度：MAE %.4f  RMSE %.4f 元/kWh"
          % (metrics["mae_slot"], metrics["rmse_slot"]))
    print("拟合度　：R²(逐时段) %.4f   R²(日均值) %.4f"
          % (metrics["r2_slot"], metrics["r2_daily"]))

    write_workbook(args.output, header, eval_dates, predicted)
    print("预测电价已保存：%s（单表「预测电价」，%d 行 x %d 列）"
          % (args.output, len(eval_dates), PERIODS_PER_DAY + 1))


if __name__ == "__main__":
    main()
