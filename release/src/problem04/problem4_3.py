"""问题 4-3：预测电价用于计划和调整，实际电价用于紧急购电费用。

运行（在项目根目录下）：
    python -m src.problem04.problem4_3

输入：
    result/problem04/prediction/pred_price_rolling.xlsx   预测电价
    附件/附件1.xlsx                                       电价基准形状
    附件/附件2.xlsx                                       实际负载和实际光伏
    附件/附件3.xlsx                                       光伏预报
    附件/附件4.xlsx                                       实际电价
    result/problem02/prediction/prediction.xlsx           负载和光伏预测

输出：
    result/problem04/result4-3.xlsx                       计划、调整、充放电、紧急购电
"""
from __future__ import annotations

import argparse
from dataclasses import replace
from datetime import datetime
import json
from pathlib import Path
from time import perf_counter

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
APPENDIX_DIR = ROOT / "附件"
RESULT_DIR = ROOT / "result" / "problem04"
PRED_PRICE_PATH = RESULT_DIR / "prediction" / "pred_price_rolling.xlsx"

from src.problem03 import problem as q3model


DEFAULT_PRICE = PRED_PRICE_PATH
DEFAULT_OUTPUT = RESULT_DIR / "result4-3.xlsx"
OUTPUT_START = q3model.OUTPUT_START
OUTPUT_END = q3model.OUTPUT_END

def load_predicted_prices(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    result = {}
    for row in rows[1:]:
        if row[0] is None:
            continue
        values = np.asarray(row[1:145], dtype=float)
        if values.shape != (144,) or not np.isfinite(values).all():
            raise ValueError(f"{row[0]} 的预测电价无效")
        result[q3model.parse_date(row[0])] = np.maximum(0.0, values)
    if (
        len(result) != 334
        or min(result) != OUTPUT_START
        or max(result) != OUTPUT_END
    ):
        raise ValueError("预测电价必须覆盖 2025-02-01 至 2025-12-31")
    return result


def load_actual_prices(path: Path) -> dict:
    """读取附件 4 的实际电价，用于 1 月储能预热。"""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    result = {}
    for row in rows[1:]:
        if row[0] is None:
            continue
        values = np.asarray(row[1:145], dtype=float)
        if values.shape != (144,) or not np.isfinite(values).all():
            raise ValueError(f"{row[0]} 的实际电价无效")
        result[q3model.parse_date(row[0])] = np.maximum(0.0, values)
    return result


def write_result4_3(output: Path, records: list) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    time_headers = [q3model.interval_label(t) for t in range(144)]

    plan_sheet = workbook.create_sheet("计划购电量")
    plan_sheet.append(["日期\\时间", *time_headers, "全天购电量", "全天购电费"])
    for result in records:
        plan_sheet.append(
            [
                datetime.combine(result.target_date, datetime.min.time()),
                *result.plan_purchase.tolist(),
                float(result.plan_purchase.sum()),
                result.plan_cost,
            ]
        )

    adjusted_sheet = workbook.create_sheet("调整购电量")
    adjusted_sheet.append(["日期\\时间", *time_headers, "全天调整购电量", "全天调整费用"])
    for result in records:
        adjusted_sheet.append(
            [
                datetime.combine(result.target_date, datetime.min.time()),
                *result.adjusted_purchase.tolist(),
                float(result.adjusted_purchase.sum()),
                result.adjustment_cost,
            ]
        )

    battery_sheet = workbook.create_sheet("充放电量")
    battery_sheet.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    for result in records:
        for block in range(6):
            lo, hi = block * 24, (block + 1) * 24
            battery_sheet.append(
                [
                    datetime.combine(result.target_date, datetime.min.time()),
                    f"{q3model.interval_label(lo).split('-')[0]}-"
                    f"{q3model.interval_label(hi - 1).split('-')[1]}",
                    float(result.charge[lo:hi].sum()),
                    float(result.discharge[lo:hi].sum()),
                    "0:00" if block == 0 else ("24:00" if block == 1 else None),
                    (
                        result.storage_initial
                        if block == 0
                        else result.storage_end[-1]
                    )
                    if block in (0, 1)
                    else None,
                ]
            )

    emergency_sheet = workbook.create_sheet("紧急购电量")
    emergency_sheet.append(["日期", "购电时间段", "购电量"])
    for result in records:
        runs = q3model.merge_emergency_runs(result.emergency)
        if not runs:
            emergency_sheet.append(
                [datetime.combine(result.target_date, datetime.min.time()), None, None]
            )
            continue
        for index, (label, amount) in enumerate(runs):
            emergency_sheet.append(
                [
                    datetime.combine(result.target_date, datetime.min.time())
                    if index == 0
                    else None,
                    label,
                    amount,
                ]
            )

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(name="宋体", size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.freeze_panes = "B2"
        worksheet.column_dimensions["A"].width = 13
        for column in range(2, worksheet.max_column + 1):
            worksheet.column_dimensions[get_column_letter(column)].width = 11
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main(argv=None) -> dict:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--price", type=Path, default=DEFAULT_PRICE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--summary-json", type=Path)
    args = parser.parse_args(argv)

    started = perf_counter()
    attachment4 = APPENDIX_DIR / "附件4.xlsx"
    predicted_price_by_date = load_predicted_prices(args.price)
    actual_price_by_date = load_actual_prices(attachment4)
    planning_price_by_date = {
        target: actual_price_by_date[target]
        for target in actual_price_by_date
        if target < OUTPUT_START
    }
    planning_price_by_date.update(predicted_price_by_date)

    inputs = q3model.load_inputs()
    actual_dates = inputs["actual_dates"]
    actual_load = inputs["actual_load"]
    actual_pv = inputs["actual_pv"]
    load_forecasts = inputs["load_forecasts"]
    pv_forecasts = inputs["pv_forecasts"]

    missing_prices = [
        target for target in actual_dates
        if target not in planning_price_by_date
    ]
    if missing_prices:
        raise ValueError(f"预测电价缺少日期：{missing_prices[:5]}")

    actual_by_date = {
        target: index for index, target in enumerate(actual_dates)
    }
    corrector = q3model.PvForecastCorrector(
        pv_forecasts,
        window=q3model.PV_SHAPE_WINDOW,
        use_level=True,
        use_shape=True,
    )
    residual_history = []
    storage = q3model.SOC_INITIAL
    records = []

    for target in actual_dates:
        corrector.fit(target)
        risk_margin = (
            q3model.segmented_risk_margin(
                np.asarray(residual_history[-q3model.RESIDUAL_WINDOW_DAYS:]),
                q3model.RISK_QUANTILE_OVERNIGHT,
                q3model.RISK_QUANTILE_DAY,
            )
            if residual_history
            else np.zeros(144)
        )
        actual_index = actual_by_date[target]
        base_load = load_forecasts[target]
        result = q3model.solve_day(
            target,
            planning_price_by_date[target],
            base_load,
            corrector,
            actual_load[actual_index],
            actual_pv[actual_index],
            storage,
            risk_margin,
        )
        actual_price = actual_price_by_date[target]
        actual_emergency_cost = float(
            5 * actual_price @ result.emergency
        )
        result = replace(
            result,
            emergency_cost=actual_emergency_cost,
        )
        storage = float(result.storage_end[-1])

        forecast_pv0 = corrector.expand(
            pv_forecasts[(target, 0)], 0, target
        )
        forecast_net0 = (base_load - forecast_pv0) * q3model.DT
        actual_net = (
            actual_load[actual_index] - actual_pv[actual_index]
        ) * q3model.DT
        residual_history.append(actual_net - forecast_net0)
        corrector.observe(target, actual_pv[actual_index])

        if q3model.OUTPUT_START <= target <= q3model.OUTPUT_END:
            records.append(result)

    write_result4_3(args.output, records)
    total_plan_kwh = sum(float(r.plan_purchase.sum()) for r in records)
    total_adjusted_kwh = sum(float(r.adjusted_purchase.sum()) for r in records)
    total_plan_cost = sum(r.plan_cost for r in records)
    total_adjustment_cost = sum(r.adjustment_cost for r in records)
    total_emergency_kwh = sum(float(r.emergency.sum()) for r in records)
    total_emergency_cost = sum(r.emergency_cost for r in records)

    print(f"求解日期：{records[0].target_date} 至 {records[-1].target_date}")
    print(f"计划购电量：{total_plan_kwh:.6f} kWh")
    print(f"调整后购电量：{total_adjusted_kwh:.6f} kWh")
    print(f"计划购电费：{total_plan_cost:.6f} 元")
    print(f"调整费用：{total_adjustment_cost:.6f} 元")
    print(f"紧急购电量：{total_emergency_kwh:.6f} kWh")
    print(f"紧急购电费：{total_emergency_cost:.6f} 元")
    print(
        "总费用："
        f"{total_plan_cost + total_adjustment_cost + total_emergency_cost:.6f} 元"
    )
    print(f"期末储电量：{storage:.6f} kWh")
    print(f"总耗时：{perf_counter() - started:.3f} 秒")
    print(f"结果已保存：{args.output}")
    summary = {
        "plan_kwh": total_plan_kwh,
        "plan_cost": total_plan_cost,
        "adjusted_kwh": total_adjusted_kwh,
        "adjustment_cost": total_adjustment_cost,
        "emergency_kwh": total_emergency_kwh,
        "emergency_cost": total_emergency_cost,
        "total_cost": (
            total_plan_cost + total_adjustment_cost + total_emergency_cost
        ),
        "storage_end": storage,
        "seconds": perf_counter() - started,
        "output": str(args.output),
    }
    if args.summary_json is not None:
        args.summary_json.parent.mkdir(parents=True, exist_ok=True)
        args.summary_json.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    return summary


if __name__ == "__main__":
    main()
