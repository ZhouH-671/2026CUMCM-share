"""问题 4-2：预测电价用于计划和计划费用，实际电价用于紧急购电费用。

运行（在项目根目录下）：
    python -m src.problem04.problem4_2

输入：
    result/problem04/prediction/pred_price_rolling.xlsx   预测电价
    result/problem02/prediction/prediction.xlsx           负载和光伏预测
    附件/附件2.xlsx                                       实际负载和实际光伏
    附件/附件4.xlsx                                       实际电价

输出：
    result/problem04/result4-2.xlsx                       计划、充放电、紧急购电
"""

from __future__ import annotations

import argparse
from datetime import datetime
import json
from pathlib import Path
from time import perf_counter

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
APPENDIX_DIR = ROOT / "附件"
RESULT_DIR = ROOT / "result" / "problem04"
PRED_PRICE_PATH = RESULT_DIR / "prediction" / "pred_price_rolling.xlsx"
PRED_LOAD_PV_PATH = ROOT / "result" / "problem02" / "prediction" / "prediction.xlsx"

from src.problem02 import problem


DEFAULT_PRICE = PRED_PRICE_PATH
DEFAULT_OUTPUT = RESULT_DIR / "result4-2.xlsx"
OUTPUT_START = problem.OUTPUT_START
OUTPUT_END = problem.OUTPUT_END

def load_predicted_prices(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    result = {}
    for row in rows[1:]:
        if row[0] is None:
            continue
        values = np.asarray(row[1:145], dtype=float)
        if values.shape != (144,) or not np.isfinite(values).all():
            raise ValueError(f"{row[0]} 的预测电价无效")
        result[problem.excel_date(row[0])] = np.maximum(0.0, values)
    if (
        len(result) != 334
        or min(result) != OUTPUT_START
        or max(result) != OUTPUT_END
    ):
        raise ValueError("预测电价必须覆盖 2025-02-01 至 2025-12-31")
    return result


def load_actual_prices(path: Path) -> dict:
    """读取附件 4 的实际电价，用于 1 月储能预热。"""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    result = {}
    for row in rows[1:]:
        if row[0] is None:
            continue
        values = np.asarray(row[1:145], dtype=float)
        if values.shape != (144,) or not np.isfinite(values).all():
            raise ValueError(f"{row[0]} 的实际电价无效")
        result[problem.excel_date(row[0])] = np.maximum(0.0, values)
    return result


def build_forecast_maps(
    attachment2: Path,
    prediction_path: Path,
) -> tuple[list, dict, dict]:
    jan_dates, jan_load, jan_pv = problem.build_january_forecasts(attachment2)
    pred_dates, pred_load, pred_pv = problem.load_prediction_series(prediction_path)

    load_map = {
        target: jan_load[index].copy()
        for index, target in enumerate(jan_dates)
    }
    pv_map = {
        target: jan_pv[index].copy()
        for index, target in enumerate(jan_dates)
    }
    load_map.update(
        {
            target: pred_load[index].copy()
            for index, target in enumerate(pred_dates)
        }
    )
    pv_map.update(
        {
            target: pred_pv[index].copy()
            for index, target in enumerate(pred_dates)
        }
    )
    return jan_dates, load_map, pv_map


def write_result4_2(
    output: Path,
    records: list[tuple],
    planning_price_by_date: dict,
) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    plan_sheet = workbook.create_sheet("计划购电量")
    plan_sheet.append(
        ["日期\\时间", *[problem.interval_label(t) for t in range(144)],
         "全天购电量", "全天购电费"]
    )
    for target, plan, actual in records:
        plan_sheet.append(
            [
                datetime.combine(target, datetime.min.time()),
                *plan.purchase_kwh.tolist(),
                float(plan.purchase_kwh.sum()),
                float(planning_price_by_date[target] @ plan.purchase_kwh),
            ]
        )

    battery_sheet = workbook.create_sheet("充放电量")
    battery_sheet.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    for target, _, actual in records:
        for block in range(6):
            lo, hi = block * 24, (block + 1) * 24
            battery_sheet.append(
                [
                    datetime.combine(target, datetime.min.time()),
                    f"{problem.interval_label(lo).split('-')[0]}-"
                    f"{problem.interval_label(hi - 1).split('-')[1]}",
                    float(actual.charge_kwh[lo:hi].sum()),
                    float(actual.discharge_kwh[lo:hi].sum()),
                    "0:00" if block == 0 else ("24:00" if block == 1 else None),
                    (
                        actual.storage_initial_kwh
                        if block == 0
                        else actual.storage_end_kwh[-1]
                    )
                    if block in (0, 1)
                    else None,
                ]
            )

    emergency_sheet = workbook.create_sheet("紧急购电量")
    emergency_sheet.append(["日期", "购电时间段", "购电量"])
    for target, _, actual in records:
        if not actual.emergency_runs:
            emergency_sheet.append(
                [datetime.combine(target, datetime.min.time()), None, None]
            )
            continue
        for index, (label, amount) in enumerate(actual.emergency_runs):
            emergency_sheet.append(
                [
                    datetime.combine(target, datetime.min.time())
                    if index == 0
                    else None,
                    label,
                    amount,
                ]
            )

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(name="宋体", size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.freeze_panes = "B2"
        worksheet.column_dimensions["A"].width = 13
        for column in range(2, worksheet.max_column + 1):
            worksheet.column_dimensions[get_column_letter(column)].width = 11
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main(argv=None) -> dict:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--price", type=Path, default=DEFAULT_PRICE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--summary-json", type=Path)
    args = parser.parse_args(argv)

    started = perf_counter()
    attachment2 = APPENDIX_DIR / "附件2.xlsx"
    attachment4 = APPENDIX_DIR / "附件4.xlsx"
    predicted_price_by_date = load_predicted_prices(args.price)
    actual_price_by_date = load_actual_prices(attachment4)
    planning_price_by_date = {
        target: actual_price_by_date[target]
        for target in actual_price_by_date
        if target < OUTPUT_START
    }
    planning_price_by_date.update(predicted_price_by_date)

    actual_dates, actual_load, actual_pv = problem.load_actual_series(attachment2)
    _, load_map, pv_map = build_forecast_maps(
        attachment2, PRED_LOAD_PV_PATH
    )
    actual_by_date = {
        target: index for index, target in enumerate(actual_dates)
    }

    residual_history = []
    storage = problem.SOC_INITIAL
    records = []
    for target in actual_dates:
        actual_index = actual_by_date[target]
        q = (
            np.quantile(
                np.asarray(residual_history[-problem.RESIDUAL_WINDOW_DAYS:]),
                problem.RISK_QUANTILE,
                axis=0,
                method="linear",
            )
            if residual_history
            else np.zeros(144)
        )
        forecast_net = (load_map[target] - pv_map[target]) * problem.DT
        price = planning_price_by_date[target]
        plan = problem.solve_plan_day(forecast_net + q, price, storage)
        actual = problem.simulate_actual_day(
            plan.purchase_kwh,
            actual_load[actual_index],
            actual_pv[actual_index],
            storage,
        )
        storage = float(actual.storage_end_kwh[-1])
        residual_history.append(
            (actual_load[actual_index] - actual_pv[actual_index]) * problem.DT
            - forecast_net
        )
        if problem.OUTPUT_START <= target <= problem.OUTPUT_END:
            records.append((target, plan, actual))

    write_result4_2(args.output, records, planning_price_by_date)
    plan_kwh = sum(float(plan.purchase_kwh.sum()) for _, plan, _ in records)
    plan_cost = sum(
        float(planning_price_by_date[target] @ plan.purchase_kwh)
        for target, plan, _ in records
    )
    emergency_kwh = sum(float(actual.emergency_kwh.sum()) for _, _, actual in records)
    emergency_cost = sum(
        float(5 * actual_price_by_date[target] @ actual.emergency_kwh)
        for target, _, actual in records
    )
    print(f"求解日期：{records[0][0]} 至 {records[-1][0]}")
    print(f"计划购电量：{plan_kwh:.6f} kWh")
    print(f"计划购电费：{plan_cost:.6f} 元")
    print(f"紧急购电量：{emergency_kwh:.6f} kWh")
    print(f"紧急购电费：{emergency_cost:.6f} 元")
    print(f"总费用：{plan_cost + emergency_cost:.6f} 元")
    print(f"期末储电量：{storage:.6f} kWh")
    print(f"总耗时：{perf_counter() - started:.3f} 秒")
    print(f"结果已保存：{args.output}")
    summary = {
        "plan_kwh": plan_kwh,
        "plan_cost": plan_cost,
        "adjusted_kwh": plan_kwh,
        "adjustment_cost": 0.0,
        "emergency_kwh": emergency_kwh,
        "emergency_cost": emergency_cost,
        "total_cost": plan_cost + emergency_cost,
        "storage_end": storage,
        "seconds": perf_counter() - started,
        "output": str(args.output),
    }
    if args.summary_json is not None:
        args.summary_json.parent.mkdir(parents=True, exist_ok=True)
        args.summary_json.write_text(
            json.dumps(summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    return summary


if __name__ == "__main__":
    main()
