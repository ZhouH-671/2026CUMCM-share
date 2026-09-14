"""使用 comparison_load.py 和 comparison_pv.py 生成 2025-02-01 至 2025-12-31 的预测表。

运行：python prediction.py
输出：../result/pred.xlsx
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
import sys

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

START_DATE = date(2025, 2, 1)
END_DATE = date(2025, 12, 31)

from comparison_load import (  # noqa: E402
    RollingRidgeLoadV4,
    find_project_root,
    load_history as load_load_history,
    select_lambda,
)
from comparison_pv import (  # noqa: E402
    ExpandingRidgePVV4,
    load_history as load_pv_history,
    original_historical_prediction,
    select_parameters,
)


def export_times(time_labels: list[str]) -> list[str]:
    """保持附件 2 的时间列格式。"""
    result = time_labels.copy()
    if result and result[-1] == "24:00":
        result[-1] = "0:00+1"
    return result


def write_workbook(
    output: Path,
    dates,
    time_headers: list[str],
    load_values: np.ndarray,
    pv_values: np.ndarray,
) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(name="宋体", size=10, bold=True)
    body_font = Font(name="宋体", size=10)
    center = Alignment(horizontal="center", vertical="center")

    datasets = (
        ("小区负载", load_values),
        ("光伏发电实际功率", pv_values),
    )
    for sheet_name, values in datasets:
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.append(["日期\\时间", *time_headers])
        for day, row in zip(dates, values):
            worksheet.append([datetime.combine(day, datetime.min.time()), *row.tolist()])

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
        for row in worksheet.iter_rows(min_row=2):
            row[0].number_format = "yyyy-mm-dd"
            row[0].alignment = center
            row[0].font = body_font
            for cell in row[1:]:
                cell.number_format = "0.0000"
                cell.font = body_font

        worksheet.freeze_panes = "B2"
        worksheet.column_dimensions["A"].width = 13
        for column in range(2, len(time_headers) + 2):
            worksheet.column_dimensions[
                openpyxl.utils.get_column_letter(column)
            ].width = 9

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> None:
    root = find_project_root(SCRIPT_DIR)
    attachment = root / "data" / "appendix02.xlsx"

    load_dates, load_time_labels, load = load_load_history(attachment)
    pv_dates, pv_time_labels, pv = load_pv_history(attachment)
    if load_dates != pv_dates or load_time_labels != pv_time_labels:
        raise ValueError("负载和光伏表的日期或时间列不一致")

    load_lambda, load_scores = select_lambda(load_dates, load)
    load_prediction = RollingRidgeLoadV4(
        load_dates, load, load_lambda
    ).historical_prediction

    pv_original = original_historical_prediction(pv)
    pv_lambda, pv_blend, _ = select_parameters(
        pv_dates, pv, pv_original
    )
    pv_ridge = ExpandingRidgePVV4(
        pv_dates, pv, pv_lambda
    ).historical_ridge
    with np.errstate(invalid="ignore"):
        pv_prediction = pv_blend * pv_original + (1 - pv_blend) * pv_ridge

    selected = [
        index
        for index, target_date in enumerate(load_dates)
        if START_DATE <= target_date <= END_DATE
    ]
    if not selected:
        raise ValueError("所选日期范围内没有预测数据")

    selected_dates = [load_dates[index] for index in selected]
    load_prediction = np.maximum(0.0, load_prediction[selected])
    pv_prediction = np.maximum(0.0, pv_prediction[selected])
    if not np.isfinite(load_prediction).all():
        raise ValueError("所选负载预测范围内存在缺失值")
    if not np.isfinite(pv_prediction).all():
        raise ValueError("所选光伏预测范围内存在缺失值")

    output = root / "result" / "pred.xlsx"
    write_workbook(
        output,
        selected_dates,
        export_times(load_time_labels),
        load_prediction,
        pv_prediction,
    )

    print(f"负载参数：lambda={load_lambda}，得分={load_scores}")
    print(f"光伏参数：lambda={pv_lambda}，原方法融合权重={pv_blend}")
    print(f"预测范围：{START_DATE} 至 {END_DATE}")
    print(f"输出行数：{len(selected_dates)} 天，时间列：{len(load_time_labels)} 个")
    print(f"表格已保存：{output}")


if __name__ == "__main__":
    main()