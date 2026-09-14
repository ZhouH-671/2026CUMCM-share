# -*- coding: utf-8 -*-

import openpyxl
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from datetime import datetime
from pathlib import Path


# =========================
# 参数
# =========================

# 当前文件位于 ./src/plot_price_trend.py
# 数据位于 ./data/appendix04.xlsx
# 图片保存到 ./result/price_trend/
BASE_DIR = Path(__file__).resolve().parent.parent
FILE = BASE_DIR / "data" / "appendix04.xlsx"
OUTPUT_DIR = BASE_DIR / "result" / "price_trend"

# 确保输出目录存在
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# =========================
# 读取电价数据
# =========================

def load_price():

    wb = openpyxl.load_workbook(
        FILE,
        data_only=True
    )

    ws = wb.active

    # 时间标题
    headers = [
        cell.value
        for cell in ws[1][1:]
    ]

    dates = []
    prices = []

    for row in ws.iter_rows(
        min_row=2,
        values_only=True
    ):

        dates.append(
            row[0]
        )

        prices.append(
            list(row[1:])
        )

    wb.close()

    return (
        dates,
        np.array(prices),
        headers
    )


# =========================
# 时间转小时
# =========================

def time_to_hour(t):

    text = str(t)

    if "+" in text:
        return 24

    if hasattr(t, "hour"):

        return (
            t.hour
            +
            t.minute/60
        )

    return 0


# =========================
# 图1：
# 2、4、6、8、10、12月13日
# 日内电价曲线
# =========================

def plot_selected_days(
        dates,
        prices,
        headers):

    target_months = [
        2,
        4,
        6,
        8,
        10,
        12
    ]

    target_dates = [
        datetime(2025, m, 13)
        for m in target_months
    ]

    hours = [
        i/6
        for i in range(144)
    ]

    plt.figure(
        figsize=(12, 6)
    )

    for target in target_dates:

        idx = dates.index(target)

        plt.plot(
            hours,
            prices[idx],
            linewidth=2,
            label=
            target.strftime(
                "%Y-%m-%d"
            )
        )

    plt.xlabel(
        "Time (hour)"
    )

    plt.ylabel(
        "Electricity Price (yuan/kWh)"
    )

    plt.title(
        "Electricity Price Curve on 13th of Every Two Months"
    )

    plt.xticks(
        np.arange(0, 25, 2)
    )

    plt.grid(
        alpha=0.3
    )

    plt.legend()

    plt.tight_layout()

    plt.savefig(
        OUTPUT_DIR / "price_daily_selected.svg"
    )

    plt.show()


# =========================
# 图2：
# 全年0、10、16、20点价格
# =========================

def plot_year_curve(
        dates,
        prices):

    # 10分钟索引

    index_dict = {

        "00:00": 0,

        "10:00": 60,

        "16:00": 96,

        "20:00": 120

    }

    plt.figure(
        figsize=(12, 6)
    )

    for label, idx in index_dict.items():

        plt.plot(
            dates,
            prices[:, idx],
            linewidth=2,
            label=label
        )

    plt.xlabel(
        "Date"
    )

    plt.ylabel(
        "Electricity Price (yuan/kWh)"
    )

    plt.title(
        "Annual Electricity Price Variation"
    )

    plt.gca().xaxis.set_major_locator(
        mdates.MonthLocator()
    )

    plt.gca().xaxis.set_major_formatter(
        mdates.DateFormatter(
            "%m"
        )
    )

    plt.grid(
        alpha=0.3
    )

    plt.legend()

    plt.tight_layout()

    plt.savefig(
        OUTPUT_DIR / "price_year_curve.svg"
    )

    plt.show()


# =========================
# 主函数
# =========================

if __name__ == "__main__":

    dates, prices, headers = load_price()

    print(
        "读取数据：",
        len(dates),
        "天",
        prices.shape
    )

    plot_selected_days(
        dates,
        prices,
        headers
    )

    plot_year_curve(
        dates,
        prices
    )

    print(
        "绘图完成，图片已保存到：",
        OUTPUT_DIR
    )