"""电价预测对比：SARIMA / Holt-Winters / Ridge 回归 / HW-Ridge 残差修正。

运行（在项目根目录下）：
    python src/comparison.py

输入：
    data/appendix04.xlsx        2025.1.1-12.31 每天 144 个 10 分钟时段电价

输出：
    result/comparison/prediction_SARIMA.xlsx
    result/comparison/prediction_HoltWinters.xlsx
    result/comparison/prediction_Regression.xlsx
    result/comparison/prediction_HW_Ridge.xlsx
    result/comparison/evaluation.xlsx
    result/comparison/r2_curve.png

说明：
    仅做电价预测与误差评估，不涉及购电策略。
    采用“日均水平 + 日内形状”分解：
        pi_{d,t} = mu_d * s_t
    其中 mu_d 用四种方法预测：
        - SARIMA(1,1,1)(1,1,1,7)
        - Holt-Winters(趋势+周季节)
        - Ridge 回归（滞后1、滞后7、星期、月份、傅里叶项）
        - HW_Ridge：Holt-Winters 基线 + Ridge 残差修正
    s_t 用过去 SHAPE_WINDOW 天同时段中位数估计。
    评估指标：MAE、RMSE、MAPE、R^2，并给出分时段 R^2 及曲线图。
"""

from __future__ import annotations

import argparse
from datetime import date, datetime
from pathlib import Path
import warnings

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from sklearn.linear_model import Ridge
from statsmodels.tsa.holtwinters import ExponentialSmoothing
from statsmodels.tsa.statespace.sarimax import SARIMAX


# ============================================================
# 路径与常量
# ============================================================
SCRIPT_DIR = Path(__file__).resolve().parent          # .../src
PROJECT_ROOT = SCRIPT_DIR.parent                      # 项目根
DATA_DIR = PROJECT_ROOT / "data"
RESULT_DIR = PROJECT_ROOT / "result" / "comparison"

PERIODS_PER_DAY = 144
OUTPUT_START = date(2025, 2, 1)
OUTPUT_END = date(2025, 12, 31)

SHAPE_WINDOW = 7          # 日内形状估计窗口（天）
MIN_HISTORY = 14          # 最少历史天数，低于此值退化为朴素预测
RIDGE_ALPHA = 1.0         # Ridge 正则强度
FOURIER_K = 2             # 年周期傅里叶项阶数
SARIMA_ORDER = (1, 1, 1)
SARIMA_SEASONAL_ORDER = (1, 1, 1, 7)

METHODS = ("SARIMA", "HoltWinters", "Regression", "HW_Ridge")
METHOD_LABELS = {
    "SARIMA": "SARIMA",
    "HoltWinters": "Holt-Winters",
    "Regression": "Ridge 回归",
    "HW_Ridge": "HW + Ridge 残差修正",
}

# 中文字体设置（若环境无中文字体，可改为英文标签）
plt.rcParams["font.sans-serif"] = ["SimHei", "Microsoft YaHei", "DejaVu Sans"]
plt.rcParams["axes.unicode_minus"] = False


# ============================================================
# 日期解析
# ============================================================
def parse_date(value) -> date:
    """把 Excel 中的日期值统一转成 datetime.date。"""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"无法识别日期：{value!r}")


# ============================================================
# 读取附件 4
# ============================================================
def load_price_series(path: Path) -> dict[date, np.ndarray]:
    """读取附件 4，返回 {日期: 长度 144 的电价向量}。"""
    if not path.exists():
        raise FileNotFoundError(f"找不到电价文件：{path}")

    print(f"[读取] 打开电价文件：{path}")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ValueError("附件 4 为空")

    result: dict[date, np.ndarray] = {}
    for row in rows[1:]:
        if row[0] in (None, ""):
            continue
        target = parse_date(row[0])
        values = np.asarray(row[1:1 + PERIODS_PER_DAY], dtype=float)
        if values.shape != (PERIODS_PER_DAY,):
            raise ValueError(
                f"{target} 的电价列数不是 {PERIODS_PER_DAY}，实际 {values.shape}"
            )
        if not np.isfinite(values).all():
            raise ValueError(f"{target} 存在缺失值")
        result[target] = np.maximum(0.0, values)

    if not result:
        raise ValueError("附件 4 中没有有效电价数据")

    dates = sorted(result.keys())
    all_values = np.array([result[d] for d in dates], dtype=float)
    print(f"[读取] 共 {len(dates)} 天，"
          f"{dates[0]} 至 {dates[-1]}")
    print(f"[读取] 电价统计：min={all_values.min():.6f}  "
          f"max={all_values.max():.6f}  "
          f"mean={all_values.mean():.6f}  "
          f"std={all_values.std():.6f}")
    return result


# ============================================================
# 分解：日均水平 + 日内形状
# ============================================================
def decompose_price(
    price_by_date: dict[date, np.ndarray],
) -> tuple[list[date], np.ndarray, np.ndarray]:
    """返回按日期排序的日期列表、日均序列、形状矩阵。"""
    dates = sorted(price_by_date.keys())
    matrix = np.array([price_by_date[d] for d in dates], dtype=float)
    daily_mean = matrix.mean(axis=1)
    safe_mean = np.maximum(daily_mean, 1e-6)
    shape_matrix = matrix / safe_mean[:, None]
    print(f"[分解] 日均电价：min={daily_mean.min():.6f}  "
          f"max={daily_mean.max():.6f}  "
          f"mean={daily_mean.mean():.6f}")
    return dates, daily_mean, shape_matrix


def estimate_shape(shape_matrix: np.ndarray, window: int = SHAPE_WINDOW) -> np.ndarray:
    """用过去 window 天同时段中位数估计日内形状，并归一化到均值为 1。"""
    if len(shape_matrix) == 0:
        return np.ones(PERIODS_PER_DAY)
    recent = shape_matrix[-window:]
    shape = np.median(recent, axis=0)
    shape = np.maximum(shape, 1e-6)
    mean = shape.mean()
    if mean <= 0:
        return np.ones(PERIODS_PER_DAY)
    return shape / mean


# ============================================================
# 三种基础日均电价预测方法
# ============================================================
def forecast_sarima(
    history: np.ndarray,
    order: tuple[int, int, int] = SARIMA_ORDER,
    seasonal_order: tuple[int, int, int, int] = SARIMA_SEASONAL_ORDER,
) -> float:
    """用 SARIMA 预测下一个日均电价。"""
    if len(history) < MIN_HISTORY:
        return float(history[-1]) if len(history) else 0.0
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            model = SARIMAX(
                history,
                order=order,
                seasonal_order=seasonal_order,
                enforce_stationarity=False,
                enforce_invertibility=False,
            )
            fit = model.fit(disp=False)
            forecast = fit.forecast(steps=1)
            value = float(forecast[0])
        if not np.isfinite(value) or value < 0:
            return float(history[-1])
        return value
    except Exception:
        return float(history[-1])


def forecast_holtwinters(
    history: np.ndarray,
    seasonal_periods: int = 7,
) -> float:
    """用 Holt-Winters 预测下一个日均电价。"""
    if len(history) < max(MIN_HISTORY, 2 * seasonal_periods):
        return float(history[-1]) if len(history) else 0.0
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            model = ExponentialSmoothing(
                history,
                trend="add",
                seasonal="add",
                seasonal_periods=seasonal_periods,
                initialization_method="estimated",
            )
            fit = model.fit(optimized=True)
            forecast = fit.forecast(1)
            value = float(forecast[0])
        if not np.isfinite(value) or value < 0:
            return float(history[-1])
        return value
    except Exception:
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                model = ExponentialSmoothing(
                    history,
                    trend="add",
                    seasonal=None,
                    initialization_method="estimated",
                )
                fit = model.fit(optimized=True)
                forecast = fit.forecast(1)
                value = float(forecast[0])
            if not np.isfinite(value) or value < 0:
                return float(history[-1])
            return value
        except Exception:
            return float(history[-1])


def _fourier_terms(target: date, k: int = FOURIER_K) -> list[float]:
    """年周期傅里叶项，使用一年中的第几天作为相位。"""
    day_of_year = target.timetuple().tm_yday
    terms = []
    for i in range(1, k + 1):
        angle = 2.0 * np.pi * i * day_of_year / 365.0
        terms.append(np.sin(angle))
        terms.append(np.cos(angle))
    return terms


def _regression_features(
    target: date,
    lag1: float,
    lag7: float,
    use_month: bool = True,
    use_fourier: bool = True,
) -> np.ndarray:
    """构造线性回归特征：滞后1、滞后7、星期、月份、傅里叶项。"""
    weekday = target.weekday()
    features = [lag1, lag7]
    features.extend(1.0 if weekday == k else 0.0 for k in range(7))
    features.append(1.0 if weekday >= 5 else 0.0)

    if use_month:
        month = target.month
        features.extend(1.0 if month == m else 0.0 for m in range(1, 13))

    if use_fourier:
        features.extend(_fourier_terms(target, k=FOURIER_K))

    return np.asarray(features, dtype=float)


def forecast_regression(
    dates: list[date],
    daily_mean: np.ndarray,
    target: date,
    alpha: float = RIDGE_ALPHA,
) -> float:
    """用 Ridge 回归预测目标日的日均电价。"""
    if len(dates) < MIN_HISTORY:
        return float(daily_mean[-1]) if len(daily_mean) else 0.0

    X, y = [], []
    for i, d in enumerate(dates):
        if i < 7:
            continue
        lag1 = daily_mean[i - 1]
        lag7 = daily_mean[i - 7]
        X.append(_regression_features(d, lag1, lag7))
        y.append(daily_mean[i])

    if len(X) < 10:
        return float(daily_mean[-1])

    X = np.asarray(X, dtype=float)
    y = np.asarray(y, dtype=float)

    model = Ridge(alpha=alpha)
    model.fit(X, y)

    lag1 = daily_mean[-1]
    lag7 = daily_mean[-7] if len(daily_mean) >= 7 else daily_mean[-1]
    target_feat = _regression_features(target, lag1, lag7).reshape(1, -1)
    value = float(model.predict(target_feat)[0])
    if not np.isfinite(value) or value < 0:
        return float(daily_mean[-1])
    return value


# ============================================================
# 第四种方法：Holt-Winters + Ridge 残差修正
# ============================================================
def _residual_features(
    target: date,
    res_lag1: float,
    res_lag7: float,
) -> np.ndarray:
    """残差修正的特征：滞后残差1、滞后残差7、星期、月份、傅里叶项。"""
    weekday = target.weekday()
    features = [res_lag1, res_lag7]
    features.extend(1.0 if weekday == k else 0.0 for k in range(7))
    features.append(1.0 if weekday >= 5 else 0.0)

    month = target.month
    features.extend(1.0 if month == m else 0.0 for m in range(1, 13))

    features.extend(_fourier_terms(target, k=FOURIER_K))
    return np.asarray(features, dtype=float)


def forecast_hw_ridge(
    dates: list[date],
    daily_mean: np.ndarray,
    target: date,
    alpha: float = RIDGE_ALPHA,
    seasonal_periods: int = 7,
) -> float:
    """Holt-Winters 基线 + Ridge 残差修正，预测目标日日均电价。

    对每个历史日 d' < target：
        用 d' 之前的数据滚动拟合 Holt-Winters，得到 mu_hat_hw(d')
        残差 r(d') = mu(d') - mu_hat_hw(d')
    然后用 Ridge 对残差建模，特征包括滞后残差、星期、月份、傅里叶项。
    最终预测 = HW(target) + Ridge(residual)(target)。
    """
    if len(dates) < max(MIN_HISTORY, 2 * seasonal_periods + 7):
        return forecast_holtwinters(daily_mean, seasonal_periods)

    # 1. 滚动计算历史 Holt-Winters 预测与残差
    residuals: list[float] = []
    residual_dates: list[date] = []
    for i in range(len(dates)):
        if i < max(MIN_HISTORY, 2 * seasonal_periods):
            continue
        history = daily_mean[:i]
        hw_pred = forecast_holtwinters(history, seasonal_periods)
        residuals.append(float(daily_mean[i] - hw_pred))
        residual_dates.append(dates[i])

    if len(residuals) < 10:
        return forecast_holtwinters(daily_mean, seasonal_periods)

    residuals = np.asarray(residuals, dtype=float)

    # 2. 用 Ridge 对残差建模
    X, y = [], []
    for j, d in enumerate(residual_dates):
        if j < 7:
            continue
        lag1 = residuals[j - 1]
        lag7 = residuals[j - 7]
        X.append(_residual_features(d, lag1, lag7))
        y.append(residuals[j])

    if len(X) < 10:
        hw_target = forecast_holtwinters(daily_mean, seasonal_periods)
        return hw_target

    X = np.asarray(X, dtype=float)
    y = np.asarray(y, dtype=float)

    ridge = Ridge(alpha=alpha)
    ridge.fit(X, y)

    # 3. 预测目标日残差
    res_lag1 = residuals[-1]
    res_lag7 = residuals[-7] if len(residuals) >= 7 else residuals[-1]
    target_feat = _residual_features(target, res_lag1, res_lag7).reshape(1, -1)
    residual_hat = float(ridge.predict(target_feat)[0])
    if not np.isfinite(residual_hat):
        residual_hat = 0.0

    # 4. 最终预测
    hw_target = forecast_holtwinters(daily_mean, seasonal_periods)
    value = hw_target + residual_hat
    if not np.isfinite(value) or value < 0:
        return hw_target
    return value


# ============================================================
# 组合预测
# ============================================================
def combine_price(mu_hat: float, shape: np.ndarray) -> np.ndarray:
    """组合日均水平与日内形状，得到 144 维预测电价。"""
    pred = float(mu_hat) * shape
    pred = np.maximum(pred, 0.0)
    return pred


# ============================================================
# 逐日滚动生成四套预测
# ============================================================
def generate_predictions(
    price_by_date: dict[date, np.ndarray],
) -> dict[str, dict[date, np.ndarray]]:
    """对 OUTPUT_START ~ OUTPUT_END 的每一天，生成四套预测。"""
    dates, daily_mean, shape_matrix = decompose_price(price_by_date)
    index_of = {d: i for i, d in enumerate(dates)}

    target_dates = [d for d in dates if OUTPUT_START <= d <= OUTPUT_END]
    total = len(target_dates)
    print(f"[预测] 目标日期：{target_dates[0]} 至 {target_dates[-1]}，"
          f"共 {total} 天")

    predictions: dict[str, dict[date, np.ndarray]] = {m: {} for m in METHODS}

    for count, target in enumerate(target_dates, start=1):
        idx = index_of[target]
        if idx == 0:
            continue

        history_dates = dates[:idx]
        history_mean = daily_mean[:idx]
        shape = estimate_shape(shape_matrix[:idx], window=SHAPE_WINDOW)

        for method in METHODS:
            if method == "SARIMA":
                mu_hat = forecast_sarima(history_mean)
            elif method == "HoltWinters":
                mu_hat = forecast_holtwinters(history_mean)
            elif method == "Regression":
                mu_hat = forecast_regression(history_dates, history_mean, target)
            else:  # HW_Ridge
                mu_hat = forecast_hw_ridge(history_dates, history_mean, target)
            predictions[method][target] = combine_price(mu_hat, shape)

        if count % 30 == 0 or count == total:
            print(f"[预测] 进度 {count}/{total}，当前目标日 {target}")

    return predictions


# ============================================================
# 写出预测文件
# ============================================================
def _time_headers() -> list[str]:
    headers = []
    for t in range(PERIODS_PER_DAY):
        minutes = t * 10
        hour, minute = divmod(minutes, 60)
        headers.append(f"{hour}:{minute:02d}")
    return headers


def write_prediction(path: Path, predictions: dict[date, np.ndarray]) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "预测电价"

    worksheet.append(["日期\\时间", *_time_headers()])
    for target in sorted(predictions.keys()):
        row = [datetime.combine(target, datetime.min.time())]
        row.extend(predictions[target].tolist())
        worksheet.append(row)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="宋体", size=10, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.freeze_panes = "B2"
    worksheet.column_dimensions["A"].width = 13
    for column in range(2, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 8

    for row in worksheet.iter_rows(min_row=2):
        row[0].number_format = "yyyy-mm-dd"
        for cell in row[1:]:
            cell.number_format = "0.0000"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    print(f"[写出] {path}（{len(predictions)} 天）")


# ============================================================
# 误差评估（含全局 R^2）
# ============================================================
def evaluate_predictions(
    predictions: dict[str, dict[date, np.ndarray]],
    price_by_date: dict[date, np.ndarray],
) -> list[dict]:
    rows = []

    all_actual = []
    for target in predictions[METHODS[0]]:
        all_actual.append(price_by_date[target])
    all_actual = np.concatenate(all_actual)
    actual_mean = float(all_actual.mean())
    ss_total = float(np.sum((all_actual - actual_mean) ** 2))

    for method, pred in predictions.items():
        maes, rmses, mapes = [], [], []
        all_pred = []
        all_true = []
        for target, pred_vec in pred.items():
            actual_vec = price_by_date[target]
            diff = pred_vec - actual_vec
            mae = float(np.mean(np.abs(diff)))
            rmse = float(np.sqrt(np.mean(diff ** 2)))
            denom = np.maximum(np.abs(actual_vec), 1e-6)
            mape = float(np.mean(np.abs(diff) / denom))
            maes.append(mae)
            rmses.append(rmse)
            mapes.append(mape)
            all_pred.append(pred_vec)
            all_true.append(actual_vec)

        all_pred = np.concatenate(all_pred)
        all_true = np.concatenate(all_true)
        ss_res = float(np.sum((all_pred - all_true) ** 2))
        r2 = 1.0 - ss_res / ss_total if ss_total > 0 else float("nan")

        rows.append(
            {
                "method": method,
                "label": METHOD_LABELS[method],
                "MAE": float(np.mean(maes)) if maes else float("nan"),
                "RMSE": float(np.mean(rmses)) if rmses else float("nan"),
                "MAPE": float(np.mean(mapes)) if mapes else float("nan"),
                "R2": r2,
                "days": len(pred),
            }
        )

    rows.sort(key=lambda r: r["RMSE"])
    return rows


# ============================================================
# 分时段 R^2
# ============================================================
def evaluate_per_period_r2(
    predictions: dict[str, dict[date, np.ndarray]],
    price_by_date: dict[date, np.ndarray],
) -> dict[str, np.ndarray]:
    """对每个时段 t 分别计算四个模型的 R^2。"""
    result: dict[str, np.ndarray] = {}

    common_dates = sorted(predictions[METHODS[0]].keys())
    actual_matrix = np.array(
        [price_by_date[d] for d in common_dates], dtype=float
    )
    actual_mean_per_period = actual_matrix.mean(axis=0)
    ss_total_per_period = np.sum(
        (actual_matrix - actual_mean_per_period[None, :]) ** 2, axis=0
    )

    for method in METHODS:
        pred_matrix = np.array(
            [predictions[method][d] for d in common_dates], dtype=float
        )
        ss_res_per_period = np.sum(
            (pred_matrix - actual_matrix) ** 2, axis=0
        )
        with np.errstate(divide="ignore", invalid="ignore"):
            r2_per_period = 1.0 - ss_res_per_period / ss_total_per_period
        r2_per_period = np.where(
            ss_total_per_period > 0, r2_per_period, np.nan
        )
        result[method] = r2_per_period

    return result


def summarize_per_period_r2(
    r2_per_period: dict[str, np.ndarray],
) -> list[dict]:
    """给出每个模型分时段 R^2 的均值、最小值、最大值。"""
    rows = []
    for method in METHODS:
        values = r2_per_period[method]
        valid = values[np.isfinite(values)]
        rows.append(
            {
                "method": method,
                "label": METHOD_LABELS[method],
                "mean": float(np.mean(valid)) if len(valid) else float("nan"),
                "min": float(np.min(valid)) if len(valid) else float("nan"),
                "max": float(np.max(valid)) if len(valid) else float("nan"),
            }
        )
    rows.sort(key=lambda r: r["mean"], reverse=True)
    return rows


# ============================================================
# 写出评估文件
# ============================================================
def write_evaluation(
    path: Path,
    rows: list[dict],
    r2_per_period: dict[str, np.ndarray],
    r2_summary: list[dict],
) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(name="宋体", size=10, bold=True)

    # Sheet 1：整体误差
    worksheet = workbook.create_sheet("误差对比")
    worksheet.append(["方法", "MAE", "RMSE", "MAPE", "R2", "评估天数"])
    for row in rows:
        worksheet.append(
            [
                row["label"],
                row["MAE"],
                row["RMSE"],
                row["MAPE"],
                row["R2"],
                row["days"],
            ]
        )
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=2):
        for cell in row[1:5]:
            cell.number_format = "0.000000"
        row[5].number_format = "0"
    worksheet.column_dimensions["A"].width = 22
    for column in range(2, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 14

    # Sheet 2：分时段 R^2
    worksheet2 = workbook.create_sheet("分时段R2")
    headers = ["时段"] + [METHOD_LABELS[m] for m in METHODS]
    worksheet2.append(headers)
    time_headers = _time_headers()
    for t in range(PERIODS_PER_DAY):
        row = [time_headers[t]]
        for method in METHODS:
            value = r2_per_period[method][t]
            row.append(float(value) if np.isfinite(value) else None)
        worksheet2.append(row)
    for cell in worksheet2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet2.iter_rows(min_row=2):
        for cell in row[1:]:
            cell.number_format = "0.000000"
    worksheet2.column_dimensions["A"].width = 12
    for column in range(2, worksheet2.max_column + 1):
        worksheet2.column_dimensions[get_column_letter(column)].width = 18

    # Sheet 3：分时段 R^2 汇总
    worksheet3 = workbook.create_sheet("R2汇总")
    worksheet3.append(["方法", "均值", "最小值", "最大值"])
    for row in r2_summary:
        worksheet3.append([row["label"], row["mean"], row["min"], row["max"]])
    for cell in worksheet3[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet3.iter_rows(min_row=2):
        for cell in row[1:]:
            cell.number_format = "0.000000"
    worksheet3.column_dimensions["A"].width = 22
    for column in range(2, worksheet3.max_column + 1):
        worksheet3.column_dimensions[get_column_letter(column)].width = 14

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    print(f"[写出] {path}")


# ============================================================
# 分时段 R^2 曲线图
# ============================================================
def plot_per_period_r2(
    path: Path,
    r2_per_period: dict[str, np.ndarray],
) -> None:
    """绘制四种模型分时段 R^2 对比曲线。"""
    x = np.arange(PERIODS_PER_DAY)
    labels = _time_headers()

    fig, ax = plt.subplots(figsize=(14, 6))

    colors = {
        "SARIMA": "#1f77b4",
        "HoltWinters": "#ff7f0e",
        "Regression": "#2ca02c",
        "HW_Ridge": "#d62728",
    }
    markers = {
        "SARIMA": "o",
        "HoltWinters": "s",
        "Regression": "^",
        "HW_Ridge": "D",
    }

    for method in METHODS:
        values = r2_per_period[method]
        ax.plot(
            x,
            values,
            label=METHOD_LABELS[method],
            color=colors[method],
            linewidth=1.5,
            marker=markers[method],
            markersize=3,
            markevery=6,
        )

    ax.axhline(0.0, color="gray", linestyle="--", linewidth=1.0)
    ax.set_xlabel("时间段", fontsize=12)
    ax.set_ylabel("R²", fontsize=12)
    ax.set_title("四种模型分时段 R² 对比", fontsize=14)
    ax.set_xticks(np.arange(0, PERIODS_PER_DAY, 6))
    ax.set_xticklabels(
        [labels[i] for i in range(0, PERIODS_PER_DAY, 6)],
        rotation=45,
        fontsize=8,
    )
    ax.grid(True, linestyle="--", alpha=0.4)
    ax.legend(fontsize=10)

    all_values = np.concatenate([
        r2_per_period[m][np.isfinite(r2_per_period[m])] for m in METHODS
    ])
    if len(all_values):
        vmin = min(-0.1, float(np.min(all_values)) - 0.05)
        vmax = max(1.05, float(np.max(all_values)) + 0.05)
        ax.set_ylim(vmin, vmax)

    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=200)
    plt.close(fig)
    print(f"[绘图] {path}")


# ============================================================
# 主函数
# ============================================================
def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data",
        type=Path,
        default=DATA_DIR / "appendix04.xlsx",
    )
    parser.add_argument(
        "--outdir",
        type=Path,
        default=RESULT_DIR,
    )
    args = parser.parse_args()

    price_by_date = load_price_series(args.data)

    predictions = generate_predictions(price_by_date)

    for method in METHODS:
        path = args.outdir / f"prediction_{method}.xlsx"
        write_prediction(path, predictions[method])

    rows = evaluate_predictions(predictions, price_by_date)
    r2_per_period = evaluate_per_period_r2(predictions, price_by_date)
    r2_summary = summarize_per_period_r2(r2_per_period)

    eval_path = args.outdir / "evaluation.xlsx"
    write_evaluation(eval_path, rows, r2_per_period, r2_summary)

    r2_curve_path = args.outdir / "r2_curve.png"
    plot_per_period_r2(r2_curve_path, r2_per_period)

    print()
    print("=" * 80)
    print("整体误差对比（按 RMSE 升序）")
    print("=" * 80)
    print(
        f"{'方法':<22s}{'MAE':>12s}{'RMSE':>12s}"
        f"{'MAPE':>12s}{'R2':>10s}{'天数':>8s}"
    )
    for row in rows:
        print(
            f"{row['label']:<22s}"
            f"{row['MAE']:>12.6f}"
            f"{row['RMSE']:>12.6f}"
            f"{row['MAPE']:>12.6f}"
            f"{row['R2']:>10.6f}"
            f"{row['days']:>8d}"
        )

    print()
    print("=" * 80)
    print("分时段 R^2 汇总（按均值降序）")
    print("=" * 80)
    print(f"{'方法':<22s}{'均值':>12s}{'最小值':>12s}{'最大值':>12s}")
    for row in r2_summary:
        print(
            f"{row['label']:<22s}"
            f"{row['mean']:>12.6f}"
            f"{row['min']:>12.6f}"
            f"{row['max']:>12.6f}"
        )

    best = rows[0]
    print()
    print(f"整体最优方案：{best['label']}"
          f"（RMSE={best['RMSE']:.6f}，"
          f"MAE={best['MAE']:.6f}，"
          f"MAPE={best['MAPE']:.6f}，"
          f"R2={best['R2']:.6f}）")

    best_per_period = r2_summary[0]
    print(f"分时段 R^2 均值最优方案：{best_per_period['label']}"
          f"（均值={best_per_period['mean']:.6f}，"
          f"最小值={best_per_period['min']:.6f}，"
          f"最大值={best_per_period['max']:.6f}）")


if __name__ == "__main__":
    main()