"""第三问：基于附件 3 光伏预报的多阶段滚动购电—储能优化。

运行（在项目根目录下）：
    python src/problem.py

输入：
    data/appendix01.xlsx        电价
    data/appendix02.xlsx        实际负载和实际光伏
    data/appendix03.xlsx        0:00、6:00、12:00、18:00 光伏预报
    data/pred.xlsx         时间序列负载预测（第二问输出）

依赖：
    depend/question_2nd.py

输出：
    result/result3.xlsx
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import sys
from time import perf_counter

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import lil_matrix


SCRIPT_DIR = Path(__file__).resolve().parent          # .../src
PROJECT_ROOT = SCRIPT_DIR.parent                      # 项目根
DATA_DIR = PROJECT_ROOT / "data"
DEPEND_DIR = PROJECT_ROOT / "depend"
RESULT_DIR = PROJECT_ROOT / "result"

if str(DEPEND_DIR) not in sys.path:
    sys.path.insert(0, str(DEPEND_DIR))

import question_2nd as q2  # noqa: E402


PERIODS_PER_DAY = 144
DT = 1 / 6
ETA_C = 0.9
ETA_D = 0.9
SOC_MIN = 1200.0
SOC_MAX = 10800.0
SOC_INITIAL = 6000.0
SOC_TARGET = 6000.0
POWER_LIMIT = 5000.0 * DT
DECISION_HOURS = (0, 6, 12, 18)
RESIDUAL_WINDOW_DAYS = 60
RISK_QUANTILE = 0.6
OUTPUT_START = date(2025, 2, 1)
OUTPUT_END = date(2025, 12, 31)


@dataclass(frozen=True)
class IntervalPlan:
    purchase_kwh: np.ndarray
    charge_kwh: np.ndarray
    discharge_kwh: np.ndarray
    storage_end_kwh: np.ndarray


@dataclass(frozen=True)
class DailyResult:
    target_date: date
    plan_purchase: np.ndarray
    adjusted_purchase: np.ndarray
    charge: np.ndarray
    discharge: np.ndarray
    emergency: np.ndarray
    spill: np.ndarray
    storage_initial: float
    storage_end: np.ndarray
    plan_cost: float
    adjustment_cost: float
    emergency_cost: float


def find_project_root(start: Path) -> Path:
    for candidate in (start, *start.parents):
        if (
            (candidate / "data" / "appendix01.xlsx").is_file()
            and (candidate / "data" / "appendix02.xlsx").is_file()
            and (candidate / "data" / "appendix03.xlsx").is_file()
        ):
            return candidate
    raise FileNotFoundError(f"从 {start} 向上未找到 data/附件1~3.xlsx")


def interval_label(period: int) -> str:
    start = period * 10
    end = (period + 1) * 10

    def clock(minutes: int) -> str:
        day, minute_of_day = divmod(minutes, 24 * 60)
        hour, minute = divmod(minute_of_day, 60)
        suffix = "+1" if day else ""
        return f"{hour}:{minute:02d}{suffix}"

    return f"{clock(start)}-{clock(end)}"


def parse_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    parts = text.split("-")
    if len(parts) != 3:
        raise ValueError(f"无法识别日期：{value!r}")
    year, month, day = map(int, parts)
    return date(year, month, day)


def load_forecast_inputs(
    attachment2: Path, pred_path: Path
) -> tuple[list[date], dict[date, np.ndarray]]:
    """返回全年日期及各日期可用于 0:00 决策的负载预测。"""
    jan_dates, jan_load, _ = q2.build_january_forecasts(attachment2)
    pred_dates, pred_load, _ = q2.load_prediction_series(pred_path)

    forecasts = {
        target: jan_load[index].copy()
        for index, target in enumerate(jan_dates)
        if target.month == 1
    }
    forecasts.update(
        {
            target: pred_load[index].copy()
            for index, target in enumerate(pred_dates)
        }
    )
    actual_dates, _, _ = q2.load_actual_series(attachment2)
    missing = [target for target in actual_dates if target not in forecasts]
    if missing:
        raise ValueError(f"缺少负载预测日期：{missing[:5]}")
    return actual_dates, forecasts


def load_pv_forecasts(
    path: Path,
) -> dict[tuple[date, int], np.ndarray]:
    """读取附件 3 的四版小时级光伏预报。"""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.worksheets[0]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if tuple(rows[0][:2]) != ("日期", "预报时刻") or len(rows[0]) != 26:
        raise ValueError("附件 3 表头格式不正确")

    result: dict[tuple[date, int], np.ndarray] = {}
    current_date = None
    for row in rows[1:]:
        if row[0] not in (None, ""):
            current_date = parse_date(row[0])
        if current_date is None or row[1] is None:
            continue
        hour, minute = map(int, str(row[1]).strip().split(":")[:2])
        if minute != 0 or hour not in DECISION_HOURS:
            raise ValueError(f"无法识别附件 3 预报时刻：{row[1]!r}")
        values = np.asarray(row[2:26], dtype=float)
        if not np.isfinite(values).all():
            raise ValueError(f"{current_date} {row[1]} 光伏预报存在缺失值")
        result[(current_date, hour)] = np.maximum(0.0, values)

    if len(result) != 365 * 4:
        raise ValueError(f"附件 3 应有 1460 组预报，实际为 {len(result)} 组")
    return result


def expand_pv_forecast(values: np.ndarray, decision_hour: int) -> np.ndarray:
    """小时预报分段常数展开到当天 144 个 10 分钟时段。"""
    result = np.zeros(PERIODS_PER_DAY)
    start = decision_hour * 6
    for horizon, value in enumerate(values, start=1):
        lo = start + 6 * (horizon - 1)
        hi = start + 6 * horizon
        if lo >= PERIODS_PER_DAY:
            break
        result[lo:min(hi, PERIODS_PER_DAY)] = value
    return result


def solve_interval_milp(
    net_demand_kwh: np.ndarray,
    price: np.ndarray,
    storage_initial: float,
    *,
    reference_purchase: np.ndarray | None = None,
) -> IntervalPlan:
    """求解计划或调整阶段的储能 MILP。"""
    n = len(net_demand_kwh)
    adjustment = reference_purchase is not None
    adjustment_blocks = {}
    if adjustment:
        adjustment_blocks = {
            "up": slice(7 * n + 1, 8 * n + 1),
            "down": slice(8 * n + 1, 9 * n + 1),
        }
    blocks = {
        "g": slice(0, n),
        "c": slice(n, 2 * n),
        "d": slice(2 * n, 3 * n),
        "S": slice(3 * n, 4 * n + 1),
        "u": slice(4 * n + 1, 5 * n + 1),
        "z": slice(5 * n + 1, 6 * n + 1),
        "short": slice(6 * n + 1, 7 * n + 1),
    }
    blocks.update(adjustment_blocks)
    size = 9 * n + 1 if adjustment else 7 * n + 1

    objective = np.zeros(size)
    if adjustment:
        objective[blocks["up"]] = 1.5 * price
        objective[blocks["down"]] = 0.5 * price
    else:
        objective[blocks["g"]] = price
    terminal_penalty = 5.0 * float(np.mean(price))
    objective[blocks["short"]] = terminal_penalty

    integrality = np.zeros(size, dtype=int)
    integrality[blocks["z"]] = 1

    lower = np.zeros(size)
    upper = np.full(size, np.inf)
    upper[blocks["c"]] = POWER_LIMIT
    upper[blocks["d"]] = POWER_LIMIT
    upper[blocks["z"]] = 1
    lower[blocks["S"]] = SOC_MIN
    upper[blocks["S"]] = SOC_MAX
    lower[blocks["S"].start] = upper[blocks["S"].start] = storage_initial

    row_count = 5 * n + 1 if adjustment else 4 * n + 1
    matrix = lil_matrix((row_count, size), dtype=float)
    row_lower = np.full(row_count, -np.inf)
    row_upper = np.zeros(row_count)

    for t in range(n):
        g = blocks["g"].start + t
        c = blocks["c"].start + t
        d = blocks["d"].start + t
        s = blocks["S"].start + t
        u = blocks["u"].start + t
        z = blocks["z"].start + t

        matrix[t, g], matrix[t, d] = 1, 1
        matrix[t, c], matrix[t, u] = -1, -1
        row_lower[t] = row_upper[t] = net_demand_kwh[t]

        row = n + t
        matrix[row, s + 1], matrix[row, s] = 1, -1
        matrix[row, c], matrix[row, d] = -ETA_C, 1 / ETA_D
        row_lower[row] = row_upper[row] = 0

        matrix[2 * n + t, c], matrix[2 * n + t, z] = 1, -POWER_LIMIT
        matrix[3 * n + t, d], matrix[3 * n + t, z] = 1, POWER_LIMIT
        row_upper[3 * n + t] = POWER_LIMIT

        if adjustment:
            up = blocks["up"].start + t
            down = blocks["down"].start + t
            matrix[4 * n + t, g] = 1
            matrix[4 * n + t, up], matrix[4 * n + t, down] = -1, 1
            row_lower[4 * n + t] = row_upper[4 * n + t] = (
                float(reference_purchase[t])
            )

    terminal_row = row_count - 1
    matrix[terminal_row, blocks["S"].stop - 1] = 1
    matrix[terminal_row, blocks["short"].start] = 1
    row_lower[terminal_row] = SOC_TARGET
    row_upper[terminal_row] = np.inf

    result = milp(
        objective,
        integrality=integrality,
        bounds=Bounds(lower, upper),
        constraints=LinearConstraint(matrix.tocsc(), row_lower, row_upper),
        options={"mip_rel_gap": 0.0},
    )
    if result.x is None or result.status != 0:
        raise RuntimeError(f"第三问区间 MILP 求解失败：{result.message}")

    return IntervalPlan(
        purchase_kwh=result.x[blocks["g"]],
        charge_kwh=result.x[blocks["c"]],
        discharge_kwh=result.x[blocks["d"]],
        storage_end_kwh=result.x[blocks["S"]][1:],
    )


def simulate_segment(
    purchase_kwh: np.ndarray,
    actual_load_kw: np.ndarray,
    actual_pv_kw: np.ndarray,
    storage_initial: float,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, float]:
    """按实际负载和实际光伏运行一个连续时段区间。"""
    n = len(purchase_kwh)
    charge = np.zeros(n)
    discharge = np.zeros(n)
    emergency = np.zeros(n)
    spill = np.zeros(n)
    storage_end = np.zeros(n)
    storage = float(storage_initial)

    for t in range(n):
        balance = (
            purchase_kwh[t]
            + actual_pv_kw[t] * DT
            - actual_load_kw[t] * DT
        )
        if balance >= 0:
            charge[t] = min(
                balance,
                POWER_LIMIT,
                max(0.0, (SOC_MAX - storage) / ETA_C),
            )
            spill[t] = max(0.0, balance - charge[t])
        else:
            discharge[t] = min(
                -balance,
                POWER_LIMIT,
                max(0.0, ETA_D * (storage - SOC_MIN)),
            )
            emergency[t] = max(0.0, -balance - discharge[t])
        storage += ETA_C * charge[t] - discharge[t] / ETA_D
        storage = min(SOC_MAX, max(SOC_MIN, storage))
        storage_end[t] = storage

    return charge, discharge, emergency, spill, storage


def updated_load_forecast(
    base_load: np.ndarray,
    actual_load: np.ndarray,
    decision_index: int,
) -> np.ndarray:
    """利用当天已发生的负载对剩余时段做水平偏差修正。"""
    if decision_index <= 0:
        return base_load.copy()
    observed_error = (
        actual_load[:decision_index] - base_load[:decision_index]
    )
    correction = float(np.mean(observed_error))
    scale = max(float(np.mean(base_load[decision_index:])), 1.0)
    correction = float(np.clip(correction, -0.2 * scale, 0.2 * scale))
    result = base_load.copy()
    result[decision_index:] = np.maximum(
        0.0, base_load[decision_index:] + correction
    )
    return result


def solve_day(
    target: date,
    price: np.ndarray,
    base_load: np.ndarray,
    pv_forecasts: dict[tuple[date, int], np.ndarray],
    actual_load: np.ndarray,
    actual_pv: np.ndarray,
    storage_initial: float,
    risk_margin: np.ndarray,
) -> DailyResult:
    """完成一天的计划、三次调整和实际运行。"""
    pv0 = expand_pv_forecast(pv_forecasts[(target, 0)], 0)
    forecast_net0 = (base_load - pv0) * DT
    plan_net = forecast_net0 + risk_margin
    plan = solve_interval_milp(plan_net, price, storage_initial)

    final_purchase = plan.purchase_kwh.copy()
    actual_charge = np.zeros(PERIODS_PER_DAY)
    actual_discharge = np.zeros(PERIODS_PER_DAY)
    actual_emergency = np.zeros(PERIODS_PER_DAY)
    actual_spill = np.zeros(PERIODS_PER_DAY)
    storage_trace = np.zeros(PERIODS_PER_DAY)
    storage = float(storage_initial)
    cursor = 0

    for decision_hour in (6, 12, 18):
        decision_index = decision_hour * 6
        segment_start = storage
        (
            charge_segment,
            discharge_segment,
            emergency_segment,
            spill_segment,
            storage,
        ) = simulate_segment(
            final_purchase[cursor:decision_index],
            actual_load[cursor:decision_index],
            actual_pv[cursor:decision_index],
            storage,
        )
        actual_charge[cursor:decision_index] = charge_segment
        actual_discharge[cursor:decision_index] = discharge_segment
        actual_emergency[cursor:decision_index] = emergency_segment
        actual_spill[cursor:decision_index] = spill_segment
        storage_trace[cursor:decision_index] = (
            segment_start
            + np.cumsum(
                ETA_C * charge_segment - discharge_segment / ETA_D
            )
        )
        cursor = decision_index

        updated_load = updated_load_forecast(
            base_load, actual_load, decision_index
        )
        pv_current = expand_pv_forecast(
            pv_forecasts[(target, decision_hour)], decision_hour
        )
        forecast_net = (updated_load - pv_current) * DT
        future_net = (
            forecast_net[decision_index:] + risk_margin[decision_index:]
        )
        adjusted = solve_interval_milp(
            future_net,
            price[decision_index:],
            storage,
            reference_purchase=plan.purchase_kwh[decision_index:],
        )
        final_purchase[decision_index:] = adjusted.purchase_kwh

    segment_start = storage
    (
        charge_segment,
        discharge_segment,
        emergency_segment,
        spill_segment,
        storage,
    ) = simulate_segment(
        final_purchase[cursor:],
        actual_load[cursor:],
        actual_pv[cursor:],
        storage,
    )
    actual_charge[cursor:] = charge_segment
    actual_discharge[cursor:] = discharge_segment
    actual_emergency[cursor:] = emergency_segment
    actual_spill[cursor:] = spill_segment
    storage_trace[cursor:] = (
        segment_start
        + np.cumsum(
            ETA_C * charge_segment - discharge_segment / ETA_D
        )
    )

    delta = final_purchase - plan.purchase_kwh
    adjustment_cost = float(
        np.sum(
            1.5 * price * np.maximum(delta, 0.0)
            + 0.5 * price * np.maximum(-delta, 0.0)
        )
    )

    return DailyResult(
        target_date=target,
        plan_purchase=plan.purchase_kwh,
        adjusted_purchase=final_purchase,
        charge=actual_charge,
        discharge=actual_discharge,
        emergency=actual_emergency,
        spill=actual_spill,
        storage_initial=float(storage_initial),
        storage_end=storage_trace,
        plan_cost=float(price @ plan.purchase_kwh),
        adjustment_cost=adjustment_cost,
        emergency_cost=float(5 * price @ actual_emergency),
    )


def merge_emergency_runs(
    emergency: np.ndarray,
) -> list[tuple[str, float]]:
    active = np.flatnonzero(emergency > 1e-6)
    runs = []
    if not len(active):
        return runs
    start = previous = int(active[0])
    for value in active[1:]:
        period = int(value)
        if period != previous + 1:
            runs.append(
                (
                    f"{interval_label(start).split('-')[0]}-"
                    f"{interval_label(previous).split('-')[1]}",
                    float(emergency[start:previous + 1].sum()),
                )
            )
            start = period
        previous = period
    runs.append(
        (
            f"{interval_label(start).split('-')[0]}-"
            f"{interval_label(previous).split('-')[1]}",
            float(emergency[start:previous + 1].sum()),
        )
    )
    return runs


def write_result3(path: Path, results: list[DailyResult], price: np.ndarray) -> None:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    time_headers = [interval_label(t) for t in range(PERIODS_PER_DAY)]

    plan_sheet = workbook.create_sheet("计划购电量")
    plan_sheet.append(["日期\\时间", *time_headers, "全天购电量", "全天购电费"])
    for result in results:
        plan_sheet.append(
            [
                datetime.combine(result.target_date, datetime.min.time()),
                *result.plan_purchase.tolist(),
                float(result.plan_purchase.sum()),
                result.plan_cost,
            ]
        )

    adjusted_sheet = workbook.create_sheet("调整购电量")
    adjusted_sheet.append(["日期\\时间", *time_headers, "全天调整购电量", "全天调整费用"])
    for result in results:
        adjusted_sheet.append(
            [
                datetime.combine(result.target_date, datetime.min.time()),
                *result.adjusted_purchase.tolist(),
                float(result.adjusted_purchase.sum()),
                result.adjustment_cost,
            ]
        )

    battery_sheet = workbook.create_sheet("充放电量")
    battery_sheet.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    for result in results:
        for block in range(6):
            lo, hi = block * 24, (block + 1) * 24
            battery_sheet.append(
                [
                    datetime.combine(result.target_date, datetime.min.time()),
                    f"{interval_label(lo).split('-')[0]}-"
                    f"{interval_label(hi - 1).split('-')[1]}",
                    float(result.charge[lo:hi].sum()),
                    float(result.discharge[lo:hi].sum()),
                    "0:00" if block == 0 else ("24:00" if block == 1 else None),
                    (
                        result.storage_initial
                        if block == 0
                        else result.storage_end[-1]
                    )
                    if block in (0, 1)
                    else None,
                ]
            )

    emergency_sheet = workbook.create_sheet("紧急购电量")
    emergency_sheet.append(["日期", "购电时间段", "购电量"])
    for result in results:
        runs = merge_emergency_runs(result.emergency)
        if not runs:
            emergency_sheet.append(
                [datetime.combine(result.target_date, datetime.min.time()), None, None]
            )
            continue
        for index, (label, amount) in enumerate(runs):
            emergency_sheet.append(
                [
                    datetime.combine(result.target_date, datetime.min.time())
                    if index == 0
                    else None,
                    label,
                    amount,
                ]
            )

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
            cell.font = Font(name="宋体", size=10, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.freeze_panes = "B2"
        worksheet.column_dimensions["A"].width = 13
        for column in range(2, worksheet.max_column + 1):
            worksheet.column_dimensions[get_column_letter(column)].width = 11
        for row in worksheet.iter_rows(min_row=2):
            row[0].number_format = "yyyy-mm-dd"
            for cell in row[1:]:
                cell.number_format = "0.0000"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--result3",
        type=Path,
        default=RESULT_DIR / "result3.xlsx",
    )
    args = parser.parse_args()

    price = q2.load_price(DATA_DIR / "appendix01.xlsx")
    actual_dates, actual_load, actual_pv = q2.load_actual_series(
        DATA_DIR / "appendix02.xlsx"
    )
    load_forecasts = load_forecast_inputs(
        DATA_DIR / "appendix02.xlsx",
        DATA_DIR / "pred.xlsx",
    )[1]
    pv_forecasts = load_pv_forecasts(DATA_DIR / "appendix03.xlsx")

    prices_by_date = {target: price for target in actual_dates}
    actual_by_date = {
        target: index for index, target in enumerate(actual_dates)
    }
    residual_history: list[np.ndarray] = []
    storage = SOC_INITIAL
    output_results = []
    started = perf_counter()

    for target in actual_dates:
        if target not in load_forecasts:
            raise ValueError(f"{target} 缺少负载预测")
        risk_margin = (
            np.quantile(
                np.asarray(residual_history[-RESIDUAL_WINDOW_DAYS:]),
                RISK_QUANTILE,
                axis=0,
                method="linear",
            )
            if residual_history
            else np.zeros(PERIODS_PER_DAY)
        )
        actual_index = actual_by_date[target]
        base_load = load_forecasts[target]
        result = solve_day(
            target,
            prices_by_date[target],
            base_load,
            pv_forecasts,
            actual_load[actual_index],
            actual_pv[actual_index],
            storage,
            risk_margin,
        )
        storage = float(result.storage_end[-1])

        forecast_pv0 = expand_pv_forecast(pv_forecasts[(target, 0)], 0)
        forecast_net0 = (base_load - forecast_pv0) * DT
        actual_net = (
            actual_load[actual_index] - actual_pv[actual_index]
        ) * DT
        residual_history.append(actual_net - forecast_net0)

        if OUTPUT_START <= target <= OUTPUT_END:
            output_results.append(result)

    if not output_results or output_results[0].target_date != OUTPUT_START:
        raise ValueError("没有生成完整的第三问结果")

    write_result3(args.result3, output_results, price)
    total_plan = sum(float(r.plan_purchase.sum()) for r in output_results)
    total_adjusted = sum(float(r.adjusted_purchase.sum()) for r in output_results)
    total_plan_cost = sum(r.plan_cost for r in output_results)
    total_adjustment_cost = sum(r.adjustment_cost for r in output_results)
    total_emergency = sum(float(r.emergency.sum()) for r in output_results)
    total_emergency_cost = sum(r.emergency_cost for r in output_results)

    print(f"求解日期：{output_results[0].target_date} 至 {output_results[-1].target_date}")
    print(f"计划购电量：{total_plan:.6f} kWh")
    print(f"调整后购电量：{total_adjusted:.6f} kWh")
    print(f"计划购电费：{total_plan_cost:.6f} 元")
    print(f"调整费用：{total_adjustment_cost:.6f} 元")
    print(f"紧急购电量：{total_emergency:.6f} kWh")
    print(f"紧急购电费：{total_emergency_cost:.6f} 元")
    print(
        "总费用："
        f"{total_plan_cost + total_adjustment_cost + total_emergency_cost:.6f} 元"
    )
    print(f"期末储电量：{storage:.6f} kWh")
    print(f"总耗时：{perf_counter() - started:.3f} 秒")
    print(f"结果已保存：{args.result3}")


if __name__ == "__main__":
    main()